"""Arma el .xlsx final calcado del Excel real del despacho -- verificado
2026-09-23 celda por celda contra el archivo real de DIABETES
('Liquidacion de impuestos 202608-Diabetes E.I.R.L..xlsx', hojas
'IGV-RENTA' y 'AF (ULTIMO)'): mismos colores exactos (naranja FFF5811E,
turquesa FF76D8E3, celeste claro FFDAEEF3), mismas formulas de Excel (no
valores ya calculados en Python -- asi si alguien corrige un numero a mano
el resto se recalcula solo), mismo detalle de Compras por tasa (18%/10.5%,
Internas/Importadas), misma regla RMT de 300 UIT, mismo par de graficos
(1 de barras 'Evolucion Anual' con 4 series, 1 de lineas 'Evolucion Mensual'
con 2 series, ubicados AL LADO de su tabla, no debajo) y mismo formato
condicional (MIN=rojo/MAX=verde por columna, barras de datos en las
columnas Total), sin lineas de cuadricula pero con bordes finos alrededor
de cada tabla de datos."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.marker import Marker
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

from .igv_renta import uit_del_anio

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
MESES_CORTO = {"01": "Enero", "02": "Febrero", "03": "Marzo", "04": "Abril", "05": "Mayo",
               "06": "Junio", "07": "Julio", "08": "Agosto", "09": "Setiembre",
               "10": "Octubre", "11": "Noviembre", "12": "Diciembre"}

# Colores EXACTOS sacados del archivo real (no aproximados).
NARANJA = "FFF5811E"
TURQUESA = "FF76D8E3"
CELESTE_CLARO = "FFDAEEF3"
RESUMEN_BG = "FFDCE6F1"
ROJO_CF = "FFC7CE"
VERDE_CF = "C6EFCE"
AMARILLO_FLAG = "FFFFFF00"
ROJO_FONT = "FFFF0000"
TEAL_TEXTO = "FF005C74"

FMT_BASE = "#,##0"
FMT_TRIBUTO = "#,##0.00"
FMT_PCT = "0.00%"
FMT_EVOL = "#,##0.00"

_BORDE_FINO = Side(style="thin", color="FFBFBFBF")
_BORDE = Border(left=_BORDE_FINO, right=_BORDE_FINO, top=_BORDE_FINO, bottom=_BORDE_FINO)


def _fmt(n):
    return round(float(n or 0), 2)


def _hex_a_argb(hex6):
    return "FF" + hex6.lstrip("#").upper()


def _franja(ws, fila, col_ini, col_fin, texto, color, blanco=True):
    for c in range(col_ini, col_fin + 1):
        celda = ws.cell(row=fila, column=c)
        celda.fill = PatternFill("solid", fgColor=color)
        celda.font = Font(bold=True, color="FFFFFFFF" if blanco else "FF000000")
    ws.cell(row=fila, column=col_ini, value=texto)
    if col_fin > col_ini:
        ws.merge_cells(start_row=fila, start_column=col_ini, end_row=fila, end_column=col_fin)
    ws.cell(row=fila, column=col_ini).alignment = Alignment(horizontal="center")


def _bordes(ws, fila_ini, fila_fin, col_ini, col_fin):
    """Lineas finas de 'caja' alrededor de cada celda de una tabla -- las
    cuadriculas de la hoja quedan apagadas (ws.sheet_view.showGridLines),
    pero las tablas necesitan verse delimitadas igual que en el Excel real."""
    for r in range(fila_ini, fila_fin + 1):
        for c in range(col_ini, col_fin + 1):
            ws.cell(row=r, column=c).border = _BORDE


def _cf_minmax(ws, rango):
    """MIN de la columna en rojo claro, MAX en verde claro -- mismo par de
    reglas 'top10 rank=1' (Resaltar el valor mas bajo/alto) que usa el
    Excel real en las columnas de Evolucion."""
    rojo = DifferentialStyle(fill=PatternFill(fgColor=ROJO_CF, bgColor=ROJO_CF, fill_type="solid"))
    verde = DifferentialStyle(fill=PatternFill(fgColor=VERDE_CF, bgColor=VERDE_CF, fill_type="solid"))
    ws.conditional_formatting.add(rango, Rule(type="top10", rank=1, bottom=True, dxf=rojo))
    ws.conditional_formatting.add(rango, Rule(type="top10", rank=1, bottom=False, dxf=verde))


def _cf_databar(ws, rango):
    ws.conditional_formatting.add(rango, DataBarRule(
        start_type="min", end_type="max", color="638EC6"))


def _agregar_logo(ws, logo_bytes, celda):
    if not logo_bytes:
        return
    try:
        img = PILImage.open(io.BytesIO(logo_bytes))
        img.thumbnail((190, 90))
        buf = io.BytesIO()
        img.convert("RGBA").save(buf, format="PNG")
        buf.seek(0)
        ws.add_image(XLImage(buf), celda)
    except Exception:
        pass


def _set_categorias_texto(chart, ws_ref_formula):
    """openpyxl set_categories() SIEMPRE arma un NumRef aunque la categoria
    sea texto (meses) -- rompe el grafico en Excel (cada mes sale como una
    serie separada en la leyenda). Fix: armar el StrRef a mano."""
    for serie in chart.series:
        serie.cat = AxDataSource(strRef=StrRef(f=ws_ref_formula))


# ---------------------------------------------------------------------
# HOJA IGV-RENTA
# ---------------------------------------------------------------------

def _hoja_igv_renta(wb, calc, branding, logo_bytes):
    ws = wb.active
    ws.title = "IGV-RENTA"
    ws.sheet_view.showGridLines = False
    v, c = calc["ventas"], calc["compras"]
    ct = c["por_tasa"]

    # OJO: D siempre queda angosta (2-3) a proposito -- es solo el
    # "colchon" izquierdo del bloque D:F que se combina para el texto de
    # descripcion en las filas de Ventas/Compras. Ninguna seccion debe
    # poner un VALOR propio en D directamente (por eso Resumen usa E para
    # el IMPORTE, no D -- si no, la columna se ve vacia/rota en las
    # secciones que no combinan D:F).
    for col, w in [("A", 2.5), ("B", 16), ("C", 18), ("D", 3), ("E", 20),
                   ("F", 16), ("G", 16), ("H", 14), ("I", 14)]:
        ws.column_dimensions[col].width = w

    def val(row, col, value=None, fmt=None, bold=False, color=None, fill=None):
        cell = ws.cell(row=row, column=col, value=value)
        if fmt:
            cell.number_format = fmt
        if bold or color:
            cell.font = Font(bold=bold, color=color)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        return cell

    # --- Encabezado ---
    val(1, 2, "RAZÓN SOCIAL:", bold=True, color=TEAL_TEXTO)
    val(1, 3, branding["nombre"], bold=True, color=TEAL_TEXTO)
    val(2, 2, "RUC:", bold=True, color=TEAL_TEXTO)
    val(2, 3, calc["ruc"], bold=True, color=TEAL_TEXTO)
    val(3, 2, "RÉGIMEN TRIBUTARIO:", bold=True, color=TEAL_TEXTO)
    val(3, 3, branding.get("regimen_tributario") or "—", bold=True, color=TEAL_TEXTO)
    _agregar_logo(ws, logo_bytes, "H1")

    val(5, 2, "PRELIQUIDACIÓN DE IMPUESTOS  MENSUAL", bold=True)
    val(6, 2, MESES_CORTO[calc["periodo"][4:]], bold=True)
    val(6, 3, calc["periodo"][:4], bold=True)

    # --- IGV CUENTA PROPIA ---
    fila = 8
    f_igv_cp_banda = fila
    _franja(ws, fila, 2, 9, "IGV CUENTA PROPIA", NARANJA)
    fila += 1
    val(fila, 2, "DESCRIPCIÓN", bold=True, color="FFFFFFFF", fill=TURQUESA)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=6)
    for cc in range(3, 7):
        ws.cell(row=fila, column=cc).fill = PatternFill("solid", fgColor=TURQUESA)
    val(fila, 7, "BASE", bold=True, color="FFFFFFFF", fill=TURQUESA)
    val(fila, 8, "TRIBUTO", bold=True, color="FFFFFFFF", fill=TURQUESA)
    val(fila, 9, "TOTAL", bold=True, color="FFFFFFFF", fill=TURQUESA)
    fila += 1

    _franja(ws, fila, 2, 9, "VENTAS", CELESTE_CLARO, blanco=False)
    fila += 1
    f_vg = fila
    val(fila, 2, "(+)", bold=True)
    val(fila, 3, "Gravadas")
    val(fila, 4, "Ventas netas")
    ws.merge_cells(start_row=fila, start_column=4, end_row=fila, end_column=6)
    val(fila, 7, _fmt(v["base_gravada"]), fmt=FMT_BASE)
    val(fila, 8, f"=ROUND(G{fila}*18%,0)", fmt=FMT_TRIBUTO)
    val(fila, 9, f"=G{fila}+H{fila}", fmt=FMT_BASE)
    fila += 1
    val(fila, 4, "Descuentos concedidos y/o devoluciones de ventas")
    ws.merge_cells(start_row=fila, start_column=4, end_row=fila, end_column=6)
    val(fila, 7, 0, fmt=FMT_BASE)
    val(fila, 8, f"=ROUND(G{fila}*18%,0)", fmt=FMT_TRIBUTO)
    val(fila, 9, f"=G{fila}+H{fila}", fmt=FMT_BASE)
    fila += 1
    f_vng = fila
    val(fila, 2, "(+)", bold=True)
    val(fila, 3, "No Gravadas")
    val(fila, 4, "Ventas netas")
    ws.merge_cells(start_row=fila, start_column=4, end_row=fila, end_column=6)
    val(fila, 7, _fmt(v["no_gravado"]), fmt=FMT_BASE)
    val(fila, 9, f"=+G{fila}", fmt=FMT_BASE)
    fila += 1
    f_vtot = fila
    val(fila, 2, "TOTAL", bold=True)
    val(fila, 7, f"=SUM(G{f_vg}:G{f_vng})", fmt=FMT_TRIBUTO, bold=True)
    val(fila, 8, f"=SUM(H{f_vg}:H{f_vng})", fmt=FMT_TRIBUTO, bold=True)
    val(fila, 9, f"=SUM(I{f_vg}:I{f_vng})", fmt=FMT_BASE, bold=True)
    _bordes(ws, f_igv_cp_banda, f_vtot, 2, 9)
    fila += 2

    f_compras_banda = fila
    _franja(ws, fila, 2, 9, "COMPRAS", CELESTE_CLARO, blanco=False)
    fila += 1

    def _fila_compra(signo, tipo, base, tasa_pct):
        nonlocal fila
        f_int = fila
        val(fila, 2, signo, bold=True)
        val(fila, 3, tipo)
        val(fila, 4, "Compras netas - Internas")
        ws.merge_cells(start_row=fila, start_column=4, end_row=fila, end_column=6)
        val(fila, 7, _fmt(base["internas"]), fmt=FMT_BASE)
        if tasa_pct is not None:
            val(fila, 8, f"=+ROUND(G{fila}*{tasa_pct}%,0)", fmt=FMT_TRIBUTO)
        val(fila, 9, f"=G{fila}+H{fila}", fmt=FMT_BASE)
        fila += 1
        f_imp = fila
        val(fila, 4, "Compras netas - Importadas")
        ws.merge_cells(start_row=fila, start_column=4, end_row=fila, end_column=6)
        val(fila, 7, _fmt(base["importadas"]), fmt=FMT_BASE)
        if tasa_pct is not None:
            val(fila, 8, f"=+ROUND(G{fila}*{tasa_pct}%,0)", fmt=FMT_TRIBUTO)
        val(fila, 9, f"=G{fila}+H{fila}", fmt=FMT_BASE)
        if not base["importadas"]:
            ws.row_dimensions[fila].hidden = True
        fila += 1
        return f_int, f_imp

    f_c18_i, _ = _fila_compra("(-)", "Gravadas 18%", ct["18"], 18)
    _fila_compra("(-)", "Gravadas 10.5%", ct["105"], 10.5)
    _, f_cng_imp = _fila_compra("(-)", "No Gravadas",
                                 {"internas": c["no_gravado"], "importadas": 0.0}, None)
    f_ctot = fila
    val(fila, 2, "TOTAL", bold=True)
    val(fila, 7, f"=SUM(G{f_c18_i}:G{f_cng_imp})", fmt=FMT_TRIBUTO, bold=True)
    val(fila, 8, f"=SUM(H{f_c18_i}:H{f_cng_imp})", fmt=FMT_TRIBUTO, bold=True)
    val(fila, 9, f"=SUM(I{f_c18_i}:I{f_cng_imp})", fmt=FMT_BASE, bold=True)
    _bordes(ws, f_compras_banda, f_ctot, 2, 9)
    fila += 2

    # --- Determinación deuda IGV ---
    f_det_igv_banda = fila
    _franja(ws, fila, 2, 9, "DETERMINACIÓN DE LA DEUDA TRIBUTARIA - IGV", NARANJA)
    fila += 1
    val(fila, 2, "DESCRIPCIÓN", bold=True, color="FFFFFFFF", fill=TURQUESA)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=8)
    for cc in range(3, 9):
        ws.cell(row=fila, column=cc).fill = PatternFill("solid", fgColor=TURQUESA)
    val(fila, 9, "TRIBUTO", bold=True, color="FFFFFFFF", fill=TURQUESA)
    fila += 1

    def _fila_i(signo, etq, formula_o_valor, negrita=False, gris=False):
        nonlocal fila
        r = fila
        val(fila, 2, signo)
        cc = val(fila, 3, etq)
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=8)
        cv = val(fila, 9, formula_o_valor, fmt=FMT_TRIBUTO)
        if negrita:
            cc.font = Font(bold=True)
            cv.font = Font(bold=True)
        if gris:
            for c2 in range(2, 10):
                ws.cell(row=fila, column=c2).fill = PatternFill("solid", fgColor=CELESTE_CLARO)
        fila += 1
        return r

    f_igv_pagar = _fila_i("(-)", f'=+IF(I{fila}>0,"Impuesto a pagar","Saldo a favor del mes")',
                           f"=+H{f_vtot}-H{f_ctot}")
    f_igv_saldoant = _fila_i("(-)", "Saldo a favor del periodo anterior",
                              _fmt(calc["igv_saldo_favor_anterior"]))
    f_igv_trib1 = fila
    _fila_i("", f'=IF(I{fila}>0,"Tributo a pagar","Saldo a favor")',
            f"=+I{f_igv_pagar}+I{f_igv_saldoant}", negrita=True, gris=True)
    f_igv_percep_per = _fila_i("(-)", "Percepciones declaradas en el periodo",
                                _fmt(calc["igv_percepciones_periodo"]))
    f_igv_percep_ant = _fila_i("(-)", "Percepciones declaradas en periodos anteriores",
                                _fmt(calc["igv_percepciones_anteriores"]))
    f_igv_percep_saldo = fila
    _fila_i("", "Saldo de percepciones no aplicadas",
            f'=IF(AND(C{f_igv_trib1}="TRIBUTO A PAGAR",I{f_igv_trib1}<ABS(I{f_igv_percep_per}+I{f_igv_percep_ant})),'
            f'I{f_igv_trib1}-ABS(I{f_igv_percep_per}+I{f_igv_percep_ant}),'
            f'IF(C{f_igv_trib1}="SALDO A FAVOR",I{f_igv_percep_per}+I{f_igv_percep_ant},0))')
    f_igv_reten_per = _fila_i("(-)", "Retenciones declaradas en el periodo",
                               _fmt(calc["igv_retenciones_periodo"]))
    f_igv_reten_ant = _fila_i("(-)", "Retenciones declaradas en periodos anteriores",
                               _fmt(calc["igv_retenciones_anteriores"]))
    _fila_i("", "Saldo de retenciones no aplicadas",
            f'=IF(ABS(I{f_igv_percep_saldo})>0,I{f_igv_reten_per}+I{f_igv_reten_ant},'
            f'IF(AND(SUM(I{f_igv_trib1}:I{f_igv_percep_ant})>0,SUM(I{f_igv_trib1}:I{f_igv_percep_ant})<ABS(I{f_igv_reten_per}+I{f_igv_reten_ant})),'
            f'SUM(I{f_igv_trib1}:I{f_igv_percep_ant})-ABS(I{f_igv_reten_per}+I{f_igv_reten_ant}),0))')
    _fila_i("(-)", "Compensación saldo a favor del exportador", 0)
    _fila_i("(-)", "Impuesto Temporal a los Activos Netos (ley n° 28424)", _fmt(calc["igv_itan"]))
    _fila_i("(-)", "Otros créditos permitido por ley", _fmt(calc["igv_otros_creditos"]))
    f_igv_totaltrib = fila
    _fila_i("", "Total tributo a pagar", f"=I{f_igv_trib1}", negrita=True, gris=True)
    _fila_i("(-)", "Pagos previos", _fmt(calc["igv_pagos_previos"]))

    _franja(ws, fila, 2, 8, "TOTAL DEUDA TRIBUTARIA DEL IGV", NARANJA)
    val(fila, 9, f"=I{f_igv_totaltrib}", bold=True, color="FFFFFFFF", fill=NARANJA)
    ws.cell(row=fila, column=9).number_format = FMT_TRIBUTO
    f_igv_total = fila
    _bordes(ws, f_det_igv_banda, f_igv_total, 2, 9)
    fila += 2

    # --- Impuesto a la Renta ---
    f_renta_cat_banda = fila
    _franja(ws, fila, 2, 9, "IMPUESTO A LA RENTA - 3ERA CATEGORIA", TURQUESA)
    fila += 1
    regla = calc["renta_regla_300_uit"]
    f_coef = fila
    val(fila, 3, "Coeficiente")
    val(fila, 5, (regla["tasa"] if regla["usa_coeficiente"] else 0), fmt=FMT_PCT)
    fila += 1
    f_pct = fila
    val(fila, 3, "Porcentaje")
    val(fila, 5, 0.01, fmt=FMT_PCT)
    fila += 1
    val(fila, 2, "DESCRIPCIÓN", bold=True, fill=RESUMEN_BG)
    val(fila, 7, "BASE", bold=True, fill=RESUMEN_BG)
    val(fila, 9, "TRIBUTO", bold=True, fill=RESUMEN_BG)
    fila += 1
    f_ingresos = fila
    val(fila, 2, "Ingresos Netos")
    val(fila, 7, f"=+G{f_vtot}", fmt=FMT_TRIBUTO)
    val(fila, 9, f"=ROUND(IF(E{f_coef}=0,G{f_ingresos}*E{f_pct},G{f_ingresos}*E{f_coef}),0)", fmt=FMT_TRIBUTO)
    _bordes(ws, f_renta_cat_banda, f_ingresos, 2, 9)
    fila += 2

    f_det_renta_banda = fila
    _franja(ws, fila, 2, 9, "DETERMINACIÓN DE LA DEUDA TRIBUTARIA - RENTA", TURQUESA)
    fila += 1
    val(fila, 2, "DESCRIPCIÓN", bold=True, color="FFFFFFFF", fill=TURQUESA)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=8)
    for cc in range(3, 9):
        ws.cell(row=fila, column=cc).fill = PatternFill("solid", fgColor=TURQUESA)
    val(fila, 9, "TRIBUTO", bold=True, color="FFFFFFFF", fill=TURQUESA)
    fila += 1

    f_renta_pagar = _fila_i("(-)", f'=+IF(I{fila}>0,"Impuesto a pagar","Saldo a favor del mes")',
                             f"=+I{f_ingresos}")
    f_renta_saldoant = fila
    saldo_ant_renta = _fmt(calc["renta_saldo_favor_anterior"])
    _fila_i("(-)", "Saldo a favor del periodo anterior", saldo_ant_renta)
    if saldo_ant_renta:
        ws.cell(row=f_renta_saldoant, column=9).fill = PatternFill("solid", fgColor=AMARILLO_FLAG)
        ws.cell(row=f_renta_saldoant, column=9).font = Font(color=ROJO_FONT)
    f_renta_trib1 = fila
    _fila_i("", f'=IF(I{fila}>0,"Tributo a pagar","Saldo a favor")',
            f"=+I{f_renta_pagar}+I{f_renta_saldoant}", negrita=True, gris=True)
    _fila_i("(-)", "Compensación saldo a favor del exportador", 0)
    _fila_i("(-)", "Impuesto Temporal a los Activos Netos (ley n° 28424)", 0)
    _fila_i("(-)", "Otros créditos permitido por ley", _fmt(calc["renta_otros_creditos"]))
    f_renta_ultima = _fila_i("(-)", "Pagos previos", _fmt(calc["renta_pagos_previos"]))

    _franja(ws, fila, 2, 8, "TOTAL DEUDA TRIBUTARIA DE RENTA DE 3ERA CATEGORIA", TURQUESA)
    val(fila, 9, f"=+IF(SUM(I{f_renta_trib1}:I{f_renta_ultima})<0,0,SUM(I{f_renta_trib1}:I{f_renta_ultima}))",
        bold=True, color="FFFFFFFF", fill=TURQUESA)
    ws.cell(row=fila, column=9).number_format = FMT_TRIBUTO
    f_renta_total = fila
    _bordes(ws, f_det_renta_banda, f_renta_total, 2, 9)
    fila += 2

    # --- Resumen -- IMPORTANTE: usa columna E para el IMPORTE (no D), D
    # queda como colchon angosto en TODA la hoja (ver comentario de anchos
    # arriba); C=codigo, D=(sin uso), E=IMPORTE, F=ESTADO, igual que el
    # Excel real (E82='IMPORTE' con columna ancha, no la D angosta).
    f_resumen_banda = fila
    _franja(ws, fila, 2, 9, "RESUMEN DEUDA TRIBUTARIA", RESUMEN_BG, blanco=False)
    fila += 1
    for c_idx, txt in [(2, "TRIBUTO"), (3, "CÓDIGO"), (5, "IMPORTE"), (6, "ESTADO")]:
        val(fila, c_idx, txt, bold=True)
    fila += 1
    f_res_igv = fila
    val(fila, 2, "IGV", bold=True)
    val(fila, 3, 1011, bold=True)
    val(fila, 5, f"=I{f_igv_total}", fmt=FMT_TRIBUTO, fill=CELESTE_CLARO)
    val(fila, 6, f'=IF(E{fila}>0,"POR PAGAR","SALDO A FAVOR")')
    fila += 1
    f_res_renta = fila
    val(fila, 2, "RENTA", bold=True)
    val(fila, 3, 3121, bold=True)
    val(fila, 5, f"=I{f_renta_total}", fmt=FMT_TRIBUTO)
    val(fila, 6, f'=IF(E{fila}>0,"POR PAGAR","SALDO A FAVOR")')
    fila += 1
    val(fila, 3, "TOTAL", bold=True)
    val(fila, 5, f"=+E{f_res_igv}+E{f_res_renta}", fmt=FMT_TRIBUTO, bold=True)
    _bordes(ws, f_resumen_banda, fila, 2, 9)

    return ws


# ---------------------------------------------------------------------
# HOJA AF (ULTIMO) -- Evolución
# ---------------------------------------------------------------------

def _hoja_evolucion(wb, branding, calc_ruc, anio_base, anio_anterior, datos_por_anio, mes_actual):
    ws = wb.create_sheet("AF (ULTIMO)")
    ws.sheet_view.showGridLines = False
    for col, w in [("A", 2.5), ("B", 15), ("C", 20), ("D", 16), ("E", 11), ("F", 11),
                   ("G", 3), ("H", 15), ("I", 15), ("J", 17), ("K", 15), ("L", 15),
                   ("M", 17), ("N", 14), ("O", 11), ("P", 3)]:
        ws.column_dimensions[col].width = w

    ws.cell(row=1, column=2, value="RAZÓN SOCIAL:").font = Font(bold=True, size=10)
    ws.cell(row=1, column=3, value=branding["nombre"]).font = Font(bold=True, size=10, color=TEAL_TEXTO)
    ws.cell(row=2, column=2, value="RUC:").font = Font(bold=True, size=10)
    ws.cell(row=2, column=3, value=calc_ruc).font = Font(bold=True, size=10)

    def sval(anio, mes_idx, clave):
        d = datos_por_anio.get(anio, [None] * 12)[mes_idx]
        return _fmt(d.get(clave)) if d else None

    anio_ant_txt = anio_anterior or ""

    fila = 4
    _franja(ws, fila, 2, 15, f"EVOLUCIÓN COMPRAS - VENTAS PERIODO {anio_ant_txt}-{anio_base}".strip("- "), NARANJA)
    fila += 2

    f_anual_banda = fila
    _franja(ws, fila, 2, 15, "EVOLUCIÓN ANUAL", NARANJA)
    fila += 1
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=6)
    ws.cell(row=fila, column=2, value="Ventas").font = Font(bold=True, color="FFFFFFFF")
    ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=13)
    ws.cell(row=fila, column=8, value="Compras").font = Font(bold=True, color="FFFFFFFF")
    for cc in list(range(2, 7)) + list(range(8, 14)):
        ws.cell(row=fila, column=cc).fill = PatternFill("solid", fgColor=TURQUESA)
    fila += 1
    f_hdr_mes = fila
    ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
    ws.cell(row=fila, column=3, value="Año").font = Font(bold=True)
    ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=6)
    ws.cell(row=fila, column=5, value="Variación").font = Font(bold=True)
    ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=10)
    ws.cell(row=fila, column=8, value=f"Año {anio_ant_txt}").font = Font(bold=True)
    ws.merge_cells(start_row=fila, start_column=11, end_row=fila, end_column=13)
    ws.cell(row=fila, column=11, value=f"Año {anio_base}").font = Font(bold=True)
    ws.merge_cells(start_row=fila, start_column=14, end_row=fila, end_column=15)
    ws.cell(row=fila, column=14, value="Variación").font = Font(bold=True)
    for cc in [3, 5, 8, 11, 14]:
        ws.cell(row=fila, column=cc).fill = PatternFill("solid", fgColor=TURQUESA)
    fila += 1
    cabeceras = {2: "Mes", 3: anio_ant_txt, 4: anio_base, 5: "Soles", 6: "%",
                 8: "Gravados", 9: "No gravados", 10: f"Total {anio_ant_txt}",
                 11: "Gravados", 12: "No Gravados", 13: f"Total {anio_base}",
                 14: "Soles", 15: "%"}
    for cc, txt in cabeceras.items():
        cel = ws.cell(row=fila, column=cc, value=txt)
        cel.font = Font(bold=True)
        cel.fill = PatternFill("solid", fgColor=TURQUESA)
    fila += 1

    f_datos_ini = fila
    for i, mes in enumerate(MESES):
        r = fila
        ws.cell(row=r, column=2, value=mes)
        ws.cell(row=r, column=3, value=sval(anio_anterior, i, "ventas") if anio_anterior else None)
        ws.cell(row=r, column=4, value=sval(anio_base, i, "ventas"))
        ws.cell(row=r, column=5, value=f'=IF(AND(C{r}=0,D{r}=0),"",+D{r}-C{r})').number_format = FMT_EVOL
        ws.cell(row=r, column=6, value=f'=IF(OR(AND(C{r}=0,D{r}=0),C{r}=0,D{r}=0),"",ABS((+D{r}/C{r})-100%))').number_format = "0.00%"
        ws.cell(row=r, column=8, value=sval(anio_anterior, i, "compras_gravadas") if anio_anterior else None)
        ws.cell(row=r, column=9, value=sval(anio_anterior, i, "compras_no_gravadas") if anio_anterior else None)
        ws.cell(row=r, column=10, value=f"=H{r}+I{r}").number_format = FMT_EVOL
        ws.cell(row=r, column=11, value=sval(anio_base, i, "compras_gravadas"))
        ws.cell(row=r, column=12, value=sval(anio_base, i, "compras_no_gravadas"))
        ws.cell(row=r, column=13, value=f"=K{r}+L{r}").number_format = FMT_EVOL
        ws.cell(row=r, column=14, value=f'=IF(AND(J{r}=0,M{r}=0),"",+M{r}-J{r})').number_format = FMT_EVOL
        ws.cell(row=r, column=15, value=f'=IF(OR(AND(J{r}=0,M{r}=0),J{r}=0,M{r}=0),"",ABS((+M{r}/J{r})-100%))').number_format = "0.00%"
        for cc in (3, 4, 8, 9, 10, 11, 12, 13):
            ws.cell(row=r, column=cc).number_format = FMT_EVOL
        fila += 1
    f_datos_fin = fila - 1

    f_total_anual = fila
    ws.cell(row=fila, column=2, value="Total General").font = Font(bold=True)
    for cc in (3, 4, 8, 9, 10, 11, 12, 13):
        letra = ws.cell(row=fila, column=cc).column_letter
        ws.cell(row=fila, column=cc, value=f"=SUBTOTAL(9,{letra}{f_datos_ini}:{letra}{f_datos_fin})")
        ws.cell(row=fila, column=cc).font = Font(bold=True)
        ws.cell(row=fila, column=cc).number_format = FMT_EVOL
    _bordes(ws, f_hdr_mes, f_total_anual, 2, 15)
    fila += 2

    # Regla 300 UIT (verificada contra el Excel real: D59=300*5500,
    # D61=IF(D34>D59,"SUPERO 1.5% O COEF","-"))
    uit = uit_del_anio(int(anio_base))
    f_uit_label = fila
    ws.cell(row=fila, column=3, value="MAX 300 UIT").font = Font(bold=True)
    ws.cell(row=fila, column=4, value=f"={300}*{uit}").number_format = FMT_EVOL
    fila += 2
    ws.cell(row=fila, column=4,
            value=f'=+IF(D{f_total_anual}>D{f_uit_label},"SUPERO 1.5% O COEF","-")')
    fila += 2

    # CF: MIN=rojo / MAX=verde por columna, barras de datos en los "Total"
    _cf_minmax(ws, f"C{f_datos_ini}:C{f_datos_fin}")
    _cf_minmax(ws, f"D{f_datos_ini}:D{f_datos_fin}")
    _cf_minmax(ws, f"H{f_datos_ini}:H{f_datos_fin}")
    _cf_minmax(ws, f"I{f_datos_ini}:I{f_datos_fin}")
    _cf_minmax(ws, f"K{f_datos_ini}:K{f_datos_fin}")
    _cf_minmax(ws, f"L{f_datos_ini}:L{f_datos_fin}")
    _cf_databar(ws, f"J{f_datos_ini}:J{f_datos_fin}")
    _cf_databar(ws, f"M{f_datos_ini}:M{f_datos_fin}")

    # --- Grafico de barras: Evolucion Anual (4 series) -- AL LADO de la
    # tabla (igual que el Excel real, cuyo grafico ancla en la columna Q),
    # no debajo.
    cat_formula = f"'AF (ULTIMO)'!$B${f_datos_ini}:$B${f_datos_fin}"
    chart_anual = BarChart()
    chart_anual.type = "col"
    chart_anual.grouping = "clustered"
    chart_anual.title = f"EVOLUCIÓN DE VENTAS - COMPRAS PERIODO: {anio_ant_txt} - {anio_base}".strip(" -")
    chart_anual.height, chart_anual.width = 10, 22
    chart_anual.y_axis.numFmt = FMT_EVOL
    chart_anual.legend.position = "b"

    color_ant = branding.get("color") or "1E3A5F"
    series_specs = [
        (3, f"VENTAS {anio_ant_txt}", "00B0F0"),
        (4, f"VENTAS {anio_base}", color_ant),
        (10, f"COMPRAS {anio_ant_txt}", "00B050"),
        (13, f"COMPRAS {anio_base}", "92D050"),
    ]
    for col, titulo, color in series_specs:
        ref = Reference(ws, min_col=col, min_row=f_datos_ini, max_row=f_datos_fin)
        chart_anual.add_data(ref, titles_from_data=False)
        s = chart_anual.series[-1]
        s.tx = SeriesLabel(v=titulo)
        s.graphicalProperties.solidFill = color
    _set_categorias_texto(chart_anual, cat_formula)
    ws.add_chart(chart_anual, f"Q{f_anual_banda}")

    # --- EVOLUCIÓN MENSUAL (mes a mes, dentro del anio base) ---
    _franja(ws, fila, 2, 15, "EVOLUCIÓN MENSUAL", NARANJA)
    f_mensual_banda = fila
    fila += 1
    ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=6)
    ws.cell(row=fila, column=3, value="Ventas").font = Font(bold=True, color="FFFFFFFF")
    ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=11)
    ws.cell(row=fila, column=8, value="Compras").font = Font(bold=True, color="FFFFFFFF")
    for cc in list(range(3, 7)) + list(range(8, 12)):
        ws.cell(row=fila, column=cc).fill = PatternFill("solid", fgColor=TURQUESA)
    fila += 1
    f_hdr2 = fila
    cabeceras2 = {2: "Mes", 3: "Soles", 4: "Variación", 6: "%",
                  8: "Gravados", 9: "No gravado", 10: "Total", 11: "Variación"}
    for cc, txt in cabeceras2.items():
        cel = ws.cell(row=fila, column=cc, value=txt)
        cel.font = Font(bold=True)
        cel.fill = PatternFill("solid", fgColor=TURQUESA)
    fila += 1

    f_m_ini = fila
    for i, mes in enumerate(MESES):
        r = fila
        r_anual = f_datos_ini + i
        ws.cell(row=r, column=2, value=mes)
        ws.cell(row=r, column=3, value=f"=D{r_anual}").number_format = FMT_EVOL
        ws.cell(row=r, column=8, value=f"=K{r_anual}").number_format = FMT_EVOL
        ws.cell(row=r, column=9, value=f"=L{r_anual}").number_format = FMT_EVOL
        ws.cell(row=r, column=10, value=f"=H{r}+I{r}").number_format = FMT_EVOL
        if i > 0:
            rp = r - 1
            ws.cell(row=r, column=4, value=f'=+IF(OR(AND(C{r}=0,C{rp}=0),C{r}=0,C{rp}=0),"",C{r}-C{rp})').number_format = FMT_EVOL
            ws.cell(row=r, column=6, value=f'=IF(OR(AND(C{r}=0,C{rp}=0),C{r}=0,C{rp}=0),"",ABS((+C{r}/C{rp})-100%))').number_format = "0.00%"
            ws.cell(row=r, column=11, value=f'=+IF(OR(AND(H{r}=0,H{rp}=0),H{r}=0,H{rp}=0),"",H{r}-H{rp})').number_format = FMT_EVOL
        fila += 1
    f_m_fin = fila - 1
    ws.cell(row=fila, column=2, value="Total General").font = Font(bold=True)
    ws.cell(row=fila, column=3, value=f"=SUM(C{f_m_ini}:C{f_m_fin})").font = Font(bold=True)
    ws.cell(row=fila, column=3).number_format = FMT_EVOL
    ws.cell(row=fila, column=8, value=f"=SUBTOTAL(9,H{f_m_ini}:H{f_m_fin})").font = Font(bold=True)
    ws.cell(row=fila, column=8).number_format = FMT_EVOL
    ws.cell(row=fila, column=10, value=f"=SUM(H{f_m_ini}:H{f_m_fin})+SUM(I{f_m_ini}:I{f_m_fin})").font = Font(bold=True)
    ws.cell(row=fila, column=10).number_format = FMT_EVOL
    _bordes(ws, f_hdr2, fila, 2, 11)

    _cf_minmax(ws, f"C{f_m_ini}:C{f_m_fin}")
    _cf_minmax(ws, f"H{f_m_ini}:H{f_m_fin}")
    _cf_minmax(ws, f"I{f_m_ini}:I{f_m_fin}")
    _cf_databar(ws, f"J{f_m_ini}:J{f_m_fin}")

    cat_formula2 = f"'AF (ULTIMO)'!$B${f_m_ini}:$B${f_m_fin}"
    chart_mensual = LineChart()
    chart_mensual.title = f"EVOLUCIÓN DE VENTAS - COMPRAS PERIODO: {anio_base}"
    chart_mensual.height, chart_mensual.width = 10, 22
    chart_mensual.y_axis.numFmt = FMT_EVOL
    chart_mensual.legend.position = "b"
    for col, titulo in [(3, f"VENTAS {anio_base}"), (10, f"COMPRAS {anio_base}")]:
        ref = Reference(ws, min_col=col, min_row=f_m_ini, max_row=f_m_fin)
        chart_mensual.add_data(ref, titles_from_data=False)
        s = chart_mensual.series[-1]
        s.tx = SeriesLabel(v=titulo)
        s.marker = Marker(symbol="circle", size=5)
        s.smooth = False
    _set_categorias_texto(chart_mensual, cat_formula2)
    ws.add_chart(chart_mensual, f"Q{f_mensual_banda}")

    return ws


def construir_workbook(ruc, calc, branding, logo_bytes, anio_base, anio_anterior, datos_por_anio):
    wb = Workbook()
    _hoja_igv_renta(wb, calc, branding, logo_bytes)
    mes_actual = int(calc["periodo"][4:6])
    _hoja_evolucion(wb, branding, ruc, anio_base, anio_anterior, datos_por_anio, mes_actual)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
